# pages/1_Pelajar_Login.py
import os, io, json, datetime, traceback
from datetime import date, datetime as dt
import pandas as pd
import streamlit as st

from lib.common import (
    ensure_db, auth_email_or_sid, get_conn, one, term_label
)

# =========================== DOCX Filler (python-docx) ===========================
from io import BytesIO
from docx import Document  # pip install python-docx

# HANYA token eksplisit: «KEY», {{KEY}}, <<KEY>>  (TIADA raw-key!)
TOKEN_FORMS = (
    lambda k: f"«{k}»",          # guillemet
    lambda k: f"{{{{{k}}}}}",    # double-curly
    lambda k: f"<<{k}>>",        # ASCII double-angle
)

def _replace_in_paragraph(p, repl: dict):
    txt = p.text or ""
    if not txt:
        return
    for k, v in repl.items():
        val = "" if v is None else str(v)
        for f in TOKEN_FORMS:
            txt = txt.replace(f(k), val)
    # rebuild runs supaya split-run Word tidak ganggu replace
    for r in list(p.runs)[::-1]:
        p._element.remove(r._element)
    p.add_run(txt)

def _replace_in_table(t, repl: dict):
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                _replace_in_paragraph(p, repl)

def _replace_in_header_footer(hf, repl: dict):
    if not hf:
        return
    for p in hf.paragraphs:
        _replace_in_paragraph(p, repl)
    for t in hf.tables:
        _replace_in_table(t, repl)

def fill_docx(template_path: str, mapping: dict) -> BytesIO:
    """Isi template DOCX menggunakan python-docx (body, tables, header/footer)."""
    doc = Document(template_path)
    for p in doc.paragraphs:
        _replace_in_paragraph(p, mapping)
    for t in doc.tables:
        _replace_in_table(t, mapping)
    for s in doc.sections:
        _replace_in_header_footer(s.header, mapping)
        _replace_in_header_footer(s.footer, mapping)
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

# ================================ App Setup =================================
st.set_page_config(page_title="Log Masuk Pelajar", page_icon="🎓", layout="wide")
st.title("Log Masuk Pelajar")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SQL_PATH = os.path.join(BASE_DIR, "..", "init_mytimes_fyp.sql")

try:
    ensure_db(SQL_PATH)
except Exception as e:
    st.error(f"Ralat DB: {e}")
    st.stop()

# ---- MIGRASI WAJIB: pastikan jadual upload wujud sebelum guna ----
def ensure_upload_tables():
    with get_conn() as conn:
        cur = conn.cursor()
        # BLI-02: jawapan industri (upload)
        cur.execute("""
            CREATE TABLE IF NOT EXISTS bli02_responses(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                term_id INTEGER,
                file_name TEXT,
                file_blob BLOB,
                uploaded_at TEXT DEFAULT (datetime('now'))
            )
        """)
        # BLI-04: bukti lapor diri (upload)
        cur.execute("""
            CREATE TABLE IF NOT EXISTS reporting_in(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                term_id INTEGER,
                reported_at TEXT,
                file_name TEXT,
                file_blob BLOB
            )
        """)
        conn.commit()

ensure_upload_tables()  # panggil awal

# ---- MIGRASI: tambah lajur week_no pada logbook jika belum ada
def ensure_logbook_has_week():
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute("PRAGMA table_info(logbook)")
        cols = {r[1] for r in cur.fetchall()}
        if "week_no" not in cols:
            cur.execute("ALTER TABLE logbook ADD COLUMN week_no INTEGER")
            conn.commit()

ensure_logbook_has_week()

# Helper: baca SQL yang kalis jadual/kolum hilang
def safe_read_sql(conn, sql, params=(), empty_cols=None):
    try:
        return pd.read_sql_query(sql, conn, params=params)
    except Exception:
        return pd.DataFrame(columns=empty_cols or [])

def require_role(roles):
    aut = st.session_state.get("auth")
    if not aut or aut.get("role_name") not in roles:
        st.error("Akses tidak dibenarkan di halaman Pelajar. Sila log masuk sebagai Pelajar.")
        if st.button("Log Keluar"):
            st.session_state.auth = None; st.rerun()
        st.stop()

# ================================ Login =====================================
if "auth" not in st.session_state:
    st.session_state.auth = None

if not st.session_state.auth:
    with st.form("login"):
        login_text = st.text_input("Emel / No. Pelajar")
        password = st.text_input("Kata Laluan", type="password")
        submitted = st.form_submit_button("Log Masuk")
    if submitted:
        user = auth_email_or_sid(login_text, password)
        if not user:
            st.info("Maklumat log masuk tidak sah.")
        elif user.get("role_name") != "student":
            st.info("Akaun ini bukan peranan Pelajar.")
        else:
            st.session_state.auth = user; st.rerun()
    st.stop()

# ============================ Selepas Login =================================
user = st.session_state.auth
require_role(["student"])
st.success(f"Log masuk sebagai {user['full_name']} ({user.get('program_code') or '-'})")

# Term semasa + tarikh LI
with get_conn() as conn:
    tlabel = term_label(conn)
    df_term = pd.read_sql_query(
        "SELECT term_id, start_date, end_date FROM terms ORDER BY term_id DESC LIMIT 1", conn
    )
if df_term.empty:
    st.warning("Tiada term aktif."); st.stop()
term_id = int(df_term.iloc[0]["term_id"])

def _fmt(d):
    try:
        return dt.strptime(d, "%Y-%m-%d").strftime("%d %B %Y") if d else ""
    except Exception:
        return ""

LI_MULA  = _fmt(df_term.iloc[0].get("start_date"))
LI_TAMAT = _fmt(df_term.iloc[0].get("end_date"))

# ================================ Metrics ===================================
with get_conn() as conn:
    bli01 = one(conn, "SELECT COUNT(1) FROM bli01 WHERE student_id=? AND term_id=?", (user['user_id'], term_id)) or 0
    bli02 = one(conn, "SELECT COUNT(1) FROM bli02_responses WHERE student_id=? AND term_id=?", (user['user_id'], term_id)) or 0
    plc   = one(conn, "SELECT COUNT(1) FROM placements WHERE student_id=? AND term_id=?", (user['user_id'], term_id)) or 0
    repin = one(conn, "SELECT COUNT(1) FROM reporting_in WHERE student_id=? AND term_id=?", (user['user_id'], term_id)) or 0
    logs  = one(conn, "SELECT COUNT(1) FROM logbook WHERE student_id=? AND term_id=?", (user['user_id'], term_id)) or 0
    rep   = one(conn, "SELECT COUNT(1) FROM final_reports WHERE student_id=? AND term_id=?", (user['user_id'], term_id)) or 0
    ind   = one(conn, "SELECT COUNT(1) FROM bli05_industry WHERE student_user_id=? AND term_id=?", (user['user_id'], term_id)) or 0
    aca   = one(conn, "SELECT COUNT(1) FROM bli08_academic WHERE student_user_id=? AND term_id=?", (user['user_id'], term_id)) or 0

c1, c2, c3 = st.columns(3)
c1.metric("Sesi", tlabel)
c2.metric("BLI-01", "✅" if bli01 else "❌")
c3.metric("BLI-02 (upload)", "✅" if bli02 else "❌")
c4, c5, c6 = st.columns(3)
c4.metric("BLI-03", "✅" if plc else "❌")
c5.metric("BLI-04 (upload)", "✅" if repin else "❌")
c6.metric("Logbook Mingguan", f"{logs} entri")
c7, c8, c9 = st.columns(3)
c7.metric("Laporan Akhir", "✅" if rep else "❌")
c8.metric("BLI-05 (Industri)", "✅" if ind else "❌")
c9.metric("BLI-08 (Akademik)", "✅" if aca else "❌")

st.divider()

# ============================ Borang & Upload ================================
st.markdown("## 📝 Borang Atas Talian & Muat Naik")
tabs = st.tabs([
    "BLI-01 Maklumat Peribadi",
    "BLI-02 (Jawapan Industri - Muat Naik)",
    "BLI-03 Pengesahan Penempatan",
    "BLI-04 (Lapor Diri - Muat Naik)"
])

# --- BLI-01 (online form)
with tabs[0]:
    with get_conn() as conn:
        df_bli01 = pd.read_sql_query(
            "SELECT data_json FROM bli01 WHERE student_id=? AND term_id=? ORDER BY id DESC LIMIT 1",
            conn, params=(user['user_id'], term_id)
        )
    data_prefill = {}
    if not df_bli01.empty and df_bli01["data_json"].iloc[0]:
        try:
            data_prefill = json.loads(df_bli01["data_json"].iloc[0]) or {}
        except Exception:
            data_prefill = {}

    with st.form("form_bli01"):
        colA, colB = st.columns(2)
        nama    = colA.text_input("Nama Penuh", value=data_prefill.get("nama", user["full_name"]))
        no_ic   = colA.text_input("No. IC", value=data_prefill.get("no_ic", ""))
        no_tel  = colA.text_input("No. Telefon", value=data_prefill.get("no_tel", ""))
        alamat  = colB.text_area("Alamat Surat-Menyurat", value=data_prefill.get("alamat", ""))
        program = colB.text_input("Kod Program", value=data_prefill.get("program", user.get("program_code") or ""))
        guardian     = st.text_input("Nama Penjaga/Waris", value=data_prefill.get("guardian", ""))
        guardian_tel = st.text_input("Telefon Penjaga/Waris", value=data_prefill.get("guardian_tel", ""))
        hantar_bli01 = st.form_submit_button("Simpan BLI-01")

    if hantar_bli01:
        payload = {
            "nama": nama, "no_ic": no_ic, "no_tel": no_tel, "alamat": alamat,
            "program": program, "guardian": guardian, "guardian_tel": guardian_tel,
            "timestamp": datetime.datetime.now().isoformat(timespec="seconds")
        }
        with get_conn() as conn:
            cur = conn.cursor()
            cur.execute("""
                INSERT INTO bli01(student_id, term_id, submitted_at, data_json)
                VALUES (?,?,datetime('now'),?)
            """, (user["user_id"], term_id, json.dumps(payload)))
            conn.commit()
        st.success("BLI-01 disimpan."); st.rerun()

# --- BLI-02 (upload jawapan industri)
with tabs[1]:
    st.caption("Muat naik jawapan/pengesahan industri (PDF/DOC/DOCX/imej).")
    uploaded = st.file_uploader("Pilih fail", type=["pdf","doc","docx","jpg","jpeg","png"], key="bli02_up")
    if uploaded and st.button("Upload BLI-02"):
        data = uploaded.read()
        with get_conn() as conn:
            cur = conn.cursor()
            cur.execute("""
                INSERT INTO bli02_responses(student_id, term_id, file_name, file_blob, uploaded_at)
                VALUES (?,?,?,?, datetime('now'))
            """, (user['user_id'], term_id, uploaded.name, data))
            conn.commit()
        st.success("BLI-02 berjaya dimuat naik."); st.rerun()

    # Senarai & muat turun terkini (kalis jadual/kolum)
    with get_conn() as conn:
        df_b2 = safe_read_sql(
            conn,
            """
            SELECT id, file_name, uploaded_at, LENGTH(file_blob) AS size
            FROM bli02_responses WHERE student_id=? AND term_id=?
            ORDER BY uploaded_at DESC
            """,
            params=(user['user_id'], term_id),
            empty_cols=["id","file_name","uploaded_at","size"]
        )
    if df_b2.empty:
        st.info("Belum ada muat naik BLI-02.")
    else:
        st.dataframe(df_b2[["file_name","uploaded_at","size"]], use_container_width=True, hide_index=True)
        with get_conn() as conn:
            cur = conn.cursor()
            cur.execute("""
                SELECT file_name, file_blob FROM bli02_responses
                WHERE student_id=? AND term_id=? ORDER BY uploaded_at DESC LIMIT 1
            """, (user['user_id'], term_id))
            row = cur.fetchone()
        if row and row[1]:
            st.download_button("⬇️ Muat Turun BLI-02 (terkini)", data=row[1],
                               file_name=row[0], type="secondary")

# --- BLI-03 (online form)
with tabs[2]:
    with get_conn() as conn:
        df_plc = pd.read_sql_query("""
            SELECT org_name, address, contact_person, contact_email, contact_phone
            FROM placements
            WHERE student_id=? AND term_id=? ORDER BY id DESC LIMIT 1
        """, conn, params=(user["user_id"], term_id))
    plc_prefill = df_plc.iloc[0].to_dict() if not df_plc.empty else {}

    with st.form("form_bli03"):
        org_name = st.text_input("Nama Organisasi", value=plc_prefill.get("org_name", ""))
        org_addr = st.text_area("Alamat Organisasi", value=plc_prefill.get("address", ""))
        contact_person = st.text_input("Penyelia Industri (Nama)", value=plc_prefill.get("contact_person", ""))
        contact_email  = st.text_input("Emel Penyelia Industri", value=plc_prefill.get("contact_email", ""))
        contact_phone  = st.text_input("Telefon Penyelia Industri", value=plc_prefill.get("contact_phone", ""))
        hantar_bli03   = st.form_submit_button("Simpan BLI-03")

    if hantar_bli03:
        if not org_name.strip():
            st.info("Nama organisasi wajib diisi.")
        else:
            with get_conn() as conn:
                cur = conn.cursor()
                cur.execute("""
                    INSERT INTO placements
                    (student_id, org_name, address, contact_person, contact_email, contact_phone, term_id, created_at)
                    VALUES (?,?,?,?,?,?,?, datetime('now'))
                """, (user["user_id"], org_name.strip(), org_addr.strip(), contact_person.strip(),
                      contact_email.strip(), contact_phone.strip(), term_id))
                conn.commit()
            st.success("BLI-03 disimpan."); st.rerun()

# --- BLI-04 (upload bukti lapor diri)
with tabs[3]:
    st.caption("Muat naik bukti Lapor Diri (PDF/DOC/DOCX/imej).")
    up_bli04 = st.file_uploader("Pilih fail", type=["pdf","doc","docx","jpg","jpeg","png"], key="bli04_up")
    if up_bli04 and st.button("Upload BLI-04"):
        blob = up_bli04.read()
        with get_conn() as conn:
            cur = conn.cursor()
            cur.execute("""
                INSERT INTO reporting_in(student_id, term_id, reported_at, file_name, file_blob)
                VALUES (?,?, datetime('now'), ?, ?)
            """, (user["user_id"], term_id, up_bli04.name, blob))
            conn.commit()
        st.success("BLI-04 berjaya dimuat naik."); st.rerun()

    with get_conn() as conn:
        df_b4 = safe_read_sql(
            conn,
            """
            SELECT id, file_name, reported_at AS uploaded_at, LENGTH(file_blob) AS size
            FROM reporting_in WHERE student_id=? AND term_id=?
            ORDER BY reported_at DESC
            """,
            params=(user['user_id'], term_id),
            empty_cols=["id","file_name","uploaded_at","size"]
        )
    if df_b4.empty:
        st.info("Belum ada muat naik BLI-04.")
    else:
        st.dataframe(df_b4[["file_name","uploaded_at","size"]], use_container_width=True, hide_index=True)
        with get_conn() as conn:
            cur = conn.cursor()
            cur.execute("""
                SELECT file_name, file_blob FROM reporting_in
                WHERE student_id=? AND term_id=? ORDER BY reported_at DESC LIMIT 1
            """, (user['user_id'], term_id))
            row = cur.fetchone()
        if row and row[1]:
            st.download_button("⬇️ Muat Turun BLI-04 (terkini)", data=row[1],
                               file_name=row[0], type="secondary")

st.divider()

# ============================ Logbook Mingguan ===============================
st.markdown("## 📒 Logbook Mingguan")

# kira minggu semasa (anggaran) dari tarikh term mula
def _suggest_week(today: date, term_start_str: str):
    try:
        ts = dt.strptime(term_start_str, "%Y-%m-%d").date()
        delta = (today - ts).days
        w = 1 + (delta // 7)
        return max(1, min(16, w))
    except Exception:
        return 1

with get_conn() as conn:
    df_logs = pd.read_sql_query("""
        SELECT log_id, week_no, entry_date, title, activities, outcomes, hours,
               COALESCE(acad_comment,'') AS acad_comment,
               COALESCE(ind_comment,'')  AS ind_comment
        FROM logbook
        WHERE student_id=? AND term_id=?
        ORDER BY COALESCE(week_no, 999), entry_date DESC
    """, conn, params=(user["user_id"], term_id))

with st.form("form_logbook"):
    c0, c1, c2 = st.columns([1,1,1])
    week_no = c0.selectbox("Minggu ke-", list(range(1, 17)),
                           index=_suggest_week(date.today(), df_term.iloc[0].get("start_date"))-1)
    entry_date = c1.date_input("Tarikh aktiviti", value=date.today())
    hours = c2.number_input("Jumlah jam (hari tersebut)", min_value=0.0, max_value=12.0, step=0.5, value=0.0)
    title = st.text_input("Ringkasan tajuk / Fokus mingguan", value="")
    activities = st.text_area("Aktiviti dilaksana (ringkas tetapi jelas)", value="", height=120)
    outcomes = st.text_area("Hasil/Output/Pembelajaran", value="", height=120)
    submit_log = st.form_submit_button("Simpan Logbook")

if submit_log:
    if not title.strip() or not activities.strip():
        st.info("Sila isi sekurang-kurangnya **Tajuk** dan **Aktiviti**.")
    else:
        with get_conn() as conn:
            cur = conn.cursor()
            # Cari entri sama tarikh (logik sedia ada), tetapi kini simpan juga week_no
            cur.execute("""
                SELECT log_id, acad_comment, ind_comment
                FROM logbook
                WHERE student_id=? AND term_id=? AND entry_date=?
            """, (user["user_id"], term_id, entry_date.isoformat()))
            row = cur.fetchone()

            def commented(r):
                return bool(r and ((r[1] and str(r[1]).strip()) or (r[2] and str(r[2]).strip())))

            if row and commented(row):
                st.info("Entri pada tarikh ini telah menerima komen penyelia dan tidak boleh diubah.")
            elif row:
                cur.execute("""
                    UPDATE logbook
                    SET week_no=?, title=?, activities=?, outcomes=?, hours=?
                    WHERE log_id=?
                """, (int(week_no), title.strip(), activities.strip(), outcomes.strip(), float(hours), int(row[0])))
                conn.commit()
                st.success("Logbook dikemas kini."); st.rerun()
            else:
                cur.execute("""
                    INSERT INTO logbook(student_id, term_id, week_no, entry_date, title, activities, outcomes, hours, created_at)
                    VALUES (?,?,?,?,?,?,?,?, datetime('now'))
                """, (user["user_id"], term_id, int(week_no), entry_date.isoformat(),
                      title.strip(), activities.strip(), outcomes.strip(), float(hours)))
                conn.commit()
                st.success("Logbook disimpan."); st.rerun()

st.markdown("### Entri Terkini")
if df_logs.empty:
    st.info("Belum ada entri logbook.")
else:
    def status_komen(r):
        a = bool(str(r["acad_comment"]).strip()); i = bool(str(r["ind_comment"]).strip())
        if a and i: return "✅ Akademik & Industri"
        if a:       return "✅ Akademik"
        if i:       return "✅ Industri"
        return "– Tiada komen"
    show = df_logs.copy()
    show["Status Komen"] = show.apply(status_komen, axis=1)
    show = show[["week_no", "entry_date", "title", "hours", "Status Komen"]]
    show = show.rename(columns={"week_no": "Minggu"})
    st.dataframe(show, use_container_width=True)

st.divider()

# =================== Surat Auto-isi (SLI-01 & SLI-03) =======================
st.markdown("## 📄 Surat Permohonan & Penempatan (Auto-isi)")
tmpl_perm = os.path.join(BASE_DIR, "..", "templates", "SLI01_Surat_Permohonan.docx")
tmpl_sli3 = os.path.join(BASE_DIR, "..", "templates", "SLI03_Surat_Penempatan.docx")

# Kumpul data profil + BLI-01 + BLI-03
with get_conn() as conn:
    u = pd.read_sql_query(
        "SELECT full_name, student_id, program_code FROM users WHERE user_id=?",
        conn, params=(user["user_id"],)
    ).iloc[0]
    df_b1 = pd.read_sql_query(
        "SELECT data_json FROM bli01 WHERE student_id=? AND term_id=? ORDER BY id DESC LIMIT 1",
        conn, params=(user["user_id"], term_id)
    )
    b1 = {}
    if not df_b1.empty and df_b1["data_json"].iloc[0]:
        try:
            b1 = json.loads(df_b1["data_json"].iloc[0]) or {}
        except Exception:
            b1 = {}
    df_p = pd.read_sql_query(
        """SELECT org_name, address, contact_person, contact_email, contact_phone
           FROM placements WHERE student_id=? AND term_id=? ORDER BY id DESC LIMIT 1""",
        conn, params=(user["user_id"], term_id)
    )
    plc_data = df_p.iloc[0].to_dict() if not df_p.empty else {}

today_str = datetime.date.today().strftime("%d %B %Y")
student_name = (b1.get("nama") or u["full_name"])
program_val  = (b1.get("program") or u["program_code"] or "")
noic_val     = (b1.get("no_ic") or "")
notel_val    = (b1.get("no_tel") or "")
alamat_val   = (b1.get("alamat") or "")
guardian_val = (b1.get("guardian") or "")
guardian_tel_val = (b1.get("guardian_tel") or "")

# ---------------------- SLI-01 ----------------------
mapping_sli01 = {
    # token legasi (chevron) & {{...}} jika ada
    "NAMA_PENUH_HURUF_BESAR": student_name.upper(),
    "NOMBOR_KAD_PENGENALAN":  noic_val,
    "NOMBOR_ID_PELAJAR":      u["student_id"] or "",
    "NAMA_PROGRAM":           program_val,
    "TARIKH_MULA_LI":         _fmt(df_term.iloc[0].get("start_date")),
    "TARIKH_TAMAT_LI":        _fmt(df_term.iloc[0].get("end_date")),
    "NAMA":                   student_name,
    "NOPELAJAR":              u["student_id"] or "",
    "PROGRAM":                program_val,
    "TARIKH":                 today_str,
    "ALAMAT":                 alamat_val,
    "NOIC":                   noic_val,
    "NOTEL":                  notel_val,
    "GUARDIAN":               guardian_val,
    "GUARDIAN_TEL":           guardian_tel_val,
    "TARIKH_SURAT":           today_str,
}

# Polisi muat turun SLI-01 (≤2 medan penting kosong)
required_fields = ["nama", "no_ic", "no_tel", "alamat", "program", "guardian", "guardian_tel"]
allowed_blanks = 2
filled = {f: bool((b1.get(f) or "").strip()) for f in required_fields}
missing = [f for f, ok in filled.items() if not ok]
can_dl_sli01 = len(missing) <= allowed_blanks
msg = f"Medan diisi: {len(required_fields)-len(missing)}/{len(required_fields)}. Boleh tinggal kosong hingga {allowed_blanks}."

MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
if not os.path.exists(tmpl_perm):
    st.error("Template SLI-01 tidak ditemui. Letak di `templates/SLI01_Surat_Permohonan.docx`.")
else:
    try:
        buf_perm = fill_docx(tmpl_perm, mapping_sli01)
        binary_doc = buf_perm.getvalue()
        if can_dl_sli01:
            st.success("SLI01: Sedia dijana. " + msg)
        else:
            st.info("SLI01: " + msg)
        st.download_button(
            "✨ Muat Turun Surat Permohonan (Auto-isi)",
            data=binary_doc,
            file_name=f"SLI01_{u['student_id']}.docx",
            mime=MIME_DOCX,
            type="secondary",
            disabled=not can_dl_sli01
        )
    except Exception:
        st.error("Gagal jana SLI-01.")
        st.code(traceback.format_exc())

# ---------------------- SLI-03 ----------------------
def _split_address_two_lines(addr: str):
    if not addr: return "", ""
    parts = [p.strip() for p in addr.replace("\n", ", ").split(",") if p.strip()]
    if len(parts) <= 1:
        return (parts[0] if parts else "", "")
    cut = max(1, len(parts)//2)
    return ", ".join(parts[:cut]), ", ".join(parts[cut:])

org_name   = plc_data.get("org_name", "")
org_addr_1, org_addr_2 = _split_address_two_lines(plc_data.get("address", ""))

# Peta ikut token <<...>> yang lazim pada templat SLI-03
mapping_sli3 = {
    # pelajar
    "NAMA_PELAJAR": student_name,
    "NO_KAD_PENGENALAN": noic_val,
    "NO_PELAJAR": u["student_id"] or "",
    "NAMA_PROGRAM": program_val,

    # tarikh LI
    "TARIKH_MULA_LATIHAN_INDUSTRI": _fmt(df_term.iloc[0].get("start_date")),
    "TARIKH_TAMAT_LATIHAN_INDUSTRI": _fmt(df_term.iloc[0].get("end_date")),

    # organisasi
    "NAMA_ORGANISASI": org_name,
    "ALAMAT_ORGANISASI_1": org_addr_1,
    "ALAMAT_ORGANISASI_2": org_addr_2,
    "BANDAR": "",
    "POSKOD": "",
    "NEGERI": "",

    # penyelaras (boleh tukar ikut kampus)
    "NAMA_PENYELARAS": "AZ’LINA BINTI ABDUL HADI",
    "JAWATAN_PENYELARAS": "Penyelaras Latihan Industri",
    "TELEFON_PENYELARAS": "012-3456789",
    "EMEL_PENYELARAS": "azlina_hadi@uitm.edu.my",
}

need_bli03 = not (org_name or "").strip()

if not os.path.exists(tmpl_sli3):
    st.error("Template SLI-03 tidak ditemui. Letak di `templates/SLI03_Surat_Penempatan.docx`.")
else:
    try:
        buf_sli3 = fill_docx(tmpl_sli3, mapping_sli3)
        if need_bli03:
            st.info("Lengkapkan **BLI-03** (Nama Organisasi) untuk auto-isi Surat Penempatan.")
        st.download_button(
            "✨ Muat Turun Surat Penempatan (Auto-isi)",
            data=buf_sli3.getvalue(),
            file_name=f"SLI03_{u['student_id']}.docx",
            mime=MIME_DOCX,
            type="secondary",
            disabled=need_bli03
        )
    except Exception:
        st.error("Gagal jana SLI-03.")
        st.code(traceback.format_exc())

# ================================ Logout ====================================
st.divider()
if st.button("Log Keluar"):
    st.session_state.auth = None; st.rerun()
